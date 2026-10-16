"""Medicaid enrollment tracker.

Run:   .venv\\Scripts\\python tracker.py              (all states)
       .venv\\Scripts\\python tracker.py Florida Kansas  (just some states)

What it does:
  1. Checks every state's website, downloading only files it doesn't already have (kept in downloads/).
  2. Pulls the total enrollment number for each month since Jan 2023.
  3. Fills gaps from manual_entries.csv (e.g. Mississippi while its site is down).
  4. Checks the results against reference.csv (reference historicals) and flags any mismatch.
  5. Saves data/enrollment.csv and writes "Medicaid Enrollment Tracker.xlsx" with MoM and QoQ growth.
"""
import sys
import traceback
from datetime import datetime

import pandas as pd

from helpers import HERE, Browser, add_months, today_key
from states import START, STATES

DATA_FILE = HERE / "data" / "enrollment.csv"
EXCEL_FILE = HERE / "Medicaid Enrollment Tracker.xlsx"


def run_states(names):
    """Run each state's reader. Returns ({state: {month: value}}, {state: error message})."""
    results, errors = {}, {}
    through = today_key()
    with Browser() as browser:
        for name in names:
            reader = STATES[name]
            print(f"  {name:<15}", end=" ", flush=True)
            try:
                needs_through = reader.__code__.co_argcount == 2
                data = reader(browser, through) if needs_through else reader(browser)
                if not data:
                    print("skipped (no automatic source - using manual_entries.csv)")
                    continue
                results[name] = data
                print(f"ok   {min(data)} to {max(data)} ({len(data)} months)")
            except Exception as e:
                errors[name] = f"{type(e).__name__}: {e}"
                print(f"FAILED  {errors[name][:120]}")
                (HERE / "data").mkdir(exist_ok=True)
                with open(HERE / "data" / "errors.log", "a", encoding="utf-8") as log:
                    log.write(f"\n=== {datetime.now():%Y-%m-%d %H:%M} {name}\n{traceback.format_exc()}")
    return results, errors


def merge(old, results):
    """Combine previous data with fresh results. Fresh numbers win; states that failed keep old numbers."""
    rows = []
    stamp = datetime.now().strftime("%Y-%m-%d %H:%M")
    for state, data in results.items():
        for month, value in data.items():
            rows.append({"state": state, "month": month, "enrollment": int(value), "source": "website", "updated": stamp})
    fresh = pd.DataFrame(rows, columns=["state", "month", "enrollment", "source", "updated"])
    kept = old[~old["state"].isin(results)] if len(old) else old
    combined = pd.concat([kept, fresh], ignore_index=True)
    combined = combined[~combined["source"].astype(str).str.startswith("manual")]  # re-applied below

    # Manual entries only fill gaps - a number from the website always wins
    manual = pd.read_csv(HERE / "manual_entries.csv", dtype={"month": str})
    have = set(zip(combined["state"], combined["month"]))
    extra = [{"state": r.state, "month": r.month, "enrollment": int(r.enrollment), "source": f"manual: {r.note}",
              "updated": stamp} for r in manual.itertuples() if (r.state, r.month) not in have]
    combined = pd.concat([combined, pd.DataFrame(extra)], ignore_index=True)
    return combined.sort_values(["state", "month"]).reset_index(drop=True)


def report_new(old, new):
    before = set(zip(old["state"], old["month"])) if len(old) else set()
    added = new[[(s, m) not in before for s, m in zip(new["state"], new["month"])]]
    if len(old) and len(added):
        print("\nNew data since last run:")
        for r in added.itertuples():
            print(f"  {r.state:<15} {r.month}  {r.enrollment:>11,}")
    elif len(old):
        print("\nNo new months since last run.")


def check_reference(data):
    """Compare against reference historicals. Any mismatch means a reader is pulling the wrong row."""
    ref = pd.read_csv(HERE / "reference.csv").set_index("state")
    lookup = {(r.state, r.month): r.enrollment for r in data.itertuples()}
    problems, checked = [], 0
    for state, row in ref.iterrows():
        for month, expected in row.items():
            got = lookup.get((state, month))
            if got is None:
                continue
            checked += 1
            if int(got) != int(expected):
                problems.append(f"  {state} {month}: tracker {got:,} vs reference {int(expected):,}")
    print(f"\nReference check: {checked - len(problems)}/{checked} match the reference historicals.")
    for p in problems:
        print(p)


def build_tables(data):
    months = []
    m = START
    while m <= data["month"].max():
        months.append(m)
        m = add_months(m, 1)
    wide = data.pivot(index="state", columns="month", values="enrollment").reindex(columns=months)
    wide = wide.reindex(sorted(STATES))
    n_states = len(wide)

    reporting = wide.notna().sum()
    total = wide.sum().where(reporting == n_states)  # only when every state has reported

    def growth(lag):
        g = wide / wide.shift(lag, axis=1) - 1
        # Total growth on a same-state basis: only states that reported both months
        both = wide.notna() & wide.shift(lag, axis=1).notna()
        now = wide.where(both).sum()
        prev = wide.shift(lag, axis=1).where(both).sum()
        g.loc["Total (same-state basis)"] = now / prev.where(prev > 0) - 1
        g.loc["States in comparison"] = both.sum()
        return g

    levels = wide.copy()
    levels.loc[f"Total (all {n_states} states)"] = total
    levels.loc["States reporting"] = reporting
    return levels, growth(1), growth(3)


def write_excel(levels, mom, qoq, errors):
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    def label(m):
        return datetime.strptime(m, "%Y-%m").strftime("%b-%y")

    header_fill = PatternFill("solid", fgColor="334A7A")
    with pd.ExcelWriter(EXCEL_FILE, engine="openpyxl") as xl:
        for name, df, fmt in [("Enrollment", levels, "#,##0"), ("MoM growth", mom, "0.0%"), ("QoQ growth", qoq, "0.0%")]:
            out = df.rename(columns=label)
            out.index.name = "State"
            out.to_excel(xl, sheet_name=name)
            ws = xl.sheets[name]
            ws.freeze_panes = "B2"
            ws.column_dimensions["A"].width = 28
            for col in range(2, ws.max_column + 1):
                ws.column_dimensions[get_column_letter(col)].width = 11
                ws.cell(1, col).alignment = Alignment(horizontal="right")
            for cell in ws[1]:
                cell.font, cell.fill = Font(bold=True, color="FFFFFF"), header_fill
            for row in ws.iter_rows(min_row=2, min_col=2):
                for cell in row:
                    is_count = ws.cell(cell.row, 1).value in ("States reporting", "States in comparison")
                    cell.number_format = "0" if is_count else fmt
            for r in range(2, ws.max_row + 1):
                if str(ws.cell(r, 1).value).startswith(("Total", "States")):
                    for c in range(1, ws.max_column + 1):
                        ws.cell(r, c).font = Font(bold=True)
        notes = [
            ["Last run", datetime.now().strftime("%Y-%m-%d %H:%M")],
            ["Total (all states)", "Only shown for months where every state has reported."],
            ["Total (same-state basis)", "Growth using only states that reported both months - comparable even when some states lag."],
            ["QoQ growth", "Change vs. three months earlier."],
            ["Minnesota", "Managed care enrollment (the series in reference historicals)."],
            ["Mississippi", "Website down; values come from manual_entries.csv."],
        ] + [[f"FAILED: {s}", e] for s, e in errors.items()]
        pd.DataFrame(notes, columns=["Item", "Note"]).to_excel(xl, sheet_name="Notes", index=False)
        xl.sheets["Notes"].column_dimensions["A"].width = 28
        xl.sheets["Notes"].column_dimensions["B"].width = 100


def main():
    names = sys.argv[1:] or list(STATES)
    unknown = [n for n in names if n not in STATES]
    if unknown:
        sys.exit(f"Unknown state(s): {unknown}. Choose from: {list(STATES)}")

    old = pd.read_csv(DATA_FILE, dtype={"month": str}) if DATA_FILE.exists() else pd.DataFrame(
        columns=["state", "month", "enrollment", "source", "updated"])
    print("Checking state websites (a Chrome window will open - leave it alone until it closes)...")
    results, errors = run_states(names)

    data = merge(old, results)
    DATA_FILE.parent.mkdir(exist_ok=True)
    data.to_csv(DATA_FILE, index=False)
    report_new(old, data)
    check_reference(data)

    try:
        write_excel(*build_tables(data), errors)
        print(f"\nSaved: {EXCEL_FILE.name}")
    except PermissionError:
        print(f"\nCould not save {EXCEL_FILE.name} - close it in Excel and run again.")
    if errors:
        print(f"\n{len(errors)} state(s) failed: {', '.join(errors)}. Details in data/errors.log")


if __name__ == "__main__":
    main()
